import re
from pathlib import Path
from typing import List
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from app.models.schemas import TestMetadata, TimeIntervalReading
from app.utils.file_utils import get_output_path

class WordService:
    @staticmethod
    def _set_cell_background(cell, color_hex: str):
        """Sets cell shading color (e.g. '1E3A8A' or 'E2E8F0')."""
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    @staticmethod
    def _set_cell_margins(cell, top=70, bottom=70, left=50, right=50):
        """Sets comfortable cell internal margins in twips (20 twips = 1 pt)."""
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    @staticmethod
    def _set_cell_border(cell, **kwargs):
        """Sets thin borders on a cell."""
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

    @classmethod
    def generate_word(
        cls,
        metadata: TestMetadata,
        intervals: List[TimeIntervalReading],
        output_filename: str = "Temperature_Rise_Test_Report.docx"
    ) -> Path:
        """
        Creates an executive-grade Word (.docx) report perfectly fitted to A4 Landscape.
        """
        doc = Document()

        # 1. Configure standard ISO A4 Landscape Page Setup (297 mm x 210 mm)
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.693)   # Exact A4 width in landscape (297 mm)
        section.page_height = Inches(8.268)   # Exact A4 height in landscape (210 mm)
        section.top_margin = Inches(0.40)
        section.bottom_margin = Inches(0.40)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)

        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(9.0)
        font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # Dynamic Component Groups from metadata.channel_labels
        raw_labels = metadata.channel_labels or []
        clean_labels = [
            str(l).strip() for l in raw_labels
            if not re.search(r'^(ambient|amb|ambt|noise|noies|sound|db|time|direction|direct|-|\s*)$', str(l).strip(), re.I)
        ]
        default_names = ["Input", "Body", "Body", "Bearing cover 1", "Bearing Cover 2", "Bearing Cover 3", "Bearing Cover 4", "Bearing Cover 5", "Output"]
        final_names = clean_labels if clean_labels else default_names
        num_channels = len(final_names)

        # Check if Noise / Vibration columns are present
        has_noise_col = any(getattr(item, 'noise', None) not in [None, ''] for item in intervals) or bool(metadata.noise_level_measured and metadata.noise_level_measured != '-')
        has_vib_col = any(getattr(item, 'vibration', None) not in [None, ''] for item in intervals)

        extra_col_count = 1 + (1 if has_noise_col else 0) + (1 if has_vib_col else 0)
        total_cols = 2 + num_channels * 2 + extra_col_count

        ambient_c = 2 + num_channels * 2
        noise_c = ambient_c + 1 if has_noise_col else None
        vib_c = (noise_c + 1) if (has_noise_col and has_vib_col) else (ambient_c + 1 if has_vib_col else None)

        # Total printable width = 11.693 - 0.90 = 10.793 inches
        fixed_w = 0.75 + 0.65 + 0.60 + (0.60 if has_noise_col else 0) + (0.65 if has_vib_col else 0)
        remaining_w = max(2.0, 10.793 - fixed_w)
        ch_w = Inches(remaining_w / (num_channels * 2)) if num_channels > 0 else Inches(0.485)

        col_widths = [Inches(0.75), Inches(0.65)]
        for _ in range(num_channels * 2):
            col_widths.append(ch_w)
        col_widths.append(Inches(0.60)) # Ambient
        if has_noise_col:
            col_widths.append(Inches(0.60)) # Noise
        if has_vib_col:
            col_widths.append(Inches(0.65)) # Vibration

        # Calculate total rows required:
        # Title (1) + Info with Noise (5) + Temp Limit (1) + Table Header (2) + Intervals (N) + Lubrication (1)
        num_rows = 1 + 5 + 1 + 2 + len(intervals) + 1
        table = doc.add_table(rows=num_rows, cols=total_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Apply column widths across all cells
        for row in table.rows:
            for c_idx, width in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = width

        def format_cell(cell, text, bold=False, color_rgb=(0x1E, 0x29, 0x3B), bg_color=None, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=9.0):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(text) if text is not None else "")
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color_rgb)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cls._set_cell_margins(cell, top=70, bottom=70, left=50, right=50)
            cls._set_cell_border(cell)
            if bg_color:
                cls._set_cell_background(cell, bg_color)

        current_row_idx = 0

        # 1. Title Banner (Row 0)
        title_cell = table.cell(current_row_idx, 0)
        for c in range(1, total_cols):
            title_cell.merge(table.cell(current_row_idx, c))
        format_cell(title_cell, metadata.test_name or "GEARBOX / MOTOR TEMPERATURE RISE TEST REPORT", bold=True, color_rgb=(0xFF, 0xFF, 0xFF), bg_color="1E3A8A", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12.0)
        current_row_idx += 1

        # 2. Info Block (Rows 1 to 4)
        serial_val = metadata.serial_number or metadata.report_number or ""
        info_pairs = [
            ("Serial / Report No:", serial_val, "Date of Test:", metadata.test_date),
            ("Product / Assembly:", metadata.product_name, "Weight:", metadata.weight),
            ("Started At:", metadata.started_at, "Direction Changed At:", metadata.direction_changed_at),
            ("Test Duration:", metadata.duration, "Conclusion:", metadata.conclusion),
        ]
        mid_split = max(1, total_cols // 2)
        for k1, v1, k2, v2 in info_pairs:
            # k1 (Cols 0 to min(2, mid_split-2))
            c_k1 = table.cell(current_row_idx, 0)
            c_k1.merge(table.cell(current_row_idx, min(2, mid_split - 2)))
            format_cell(c_k1, k1, bold=True, bg_color="E2E8F0", font_size=8.0)

            # v1 (Cols min(3, mid_split-1) to mid_split-1)
            c_v1 = table.cell(current_row_idx, min(3, mid_split - 1))
            c_v1.merge(table.cell(current_row_idx, mid_split - 1))
            format_cell(c_v1, v1, bold=False, font_size=8.0)

            # k2 (Cols mid_split to min(mid_split+2, total_cols-2))
            c_k2 = table.cell(current_row_idx, mid_split)
            c_k2.merge(table.cell(current_row_idx, min(mid_split + 2, total_cols - 2)))
            format_cell(c_k2, k2, bold=True, bg_color="E2E8F0", font_size=8.0)

            # v2 (Cols min(mid_split+3, total_cols-1) to total_cols-1)
            c_v2 = table.cell(current_row_idx, min(mid_split + 3, total_cols - 1))
            c_v2.merge(table.cell(current_row_idx, total_cols - 1))
            format_cell(c_v2, v2, bold=False, font_size=8.0)

            current_row_idx += 1

        # 3. Noise Level (Row 5)
        c_n1 = table.cell(current_row_idx, 0)
        c_n1.merge(table.cell(current_row_idx, min(2, mid_split - 2)))
        format_cell(c_n1, "Noise level", bold=True, font_size=8.0)

        c_n2 = table.cell(current_row_idx, min(3, mid_split - 1))
        c_n2.merge(table.cell(current_row_idx, mid_split - 1))
        format_cell(c_n2, metadata.noise_level_limit or "< 85 dB", font_size=8.0)

        c_n3 = table.cell(current_row_idx, mid_split)
        c_n3.merge(table.cell(current_row_idx, total_cols - 1))
        format_cell(c_n3, metadata.noise_level_measured or "-", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=8.0)
        current_row_idx += 1

        # 4. Temp Rise Limit (Row 6)
        c_tlim = table.cell(current_row_idx, 0)
        for c in range(1, total_cols):
            c_tlim.merge(table.cell(current_row_idx, c))
        format_cell(c_tlim, f"Temperature rise {metadata.temp_rise_limit or '< 40°C over the ambient ( after 1hour )'}", bold=True, bg_color="E2E8F0", font_size=8.0)
        current_row_idx += 1

        # 4. Two-Tier Headers (Rows 7 & 8)
        head_row1 = current_row_idx
        head_row2 = current_row_idx + 1

        # Time (Col 0)
        c_time_h = table.cell(head_row1, 0)
        c_time_h.merge(table.cell(head_row2, 0))
        format_cell(c_time_h, "Time", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        # Direction (Col 1)
        c_dir_h = table.cell(head_row1, 1)
        c_dir_h.merge(table.cell(head_row2, 1))
        format_cell(c_dir_h, "Direction", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        components = [
            (name, 2 + i * 2, 3 + i * 2)
            for i, name in enumerate(final_names)
        ]
        for comp_name, start_c, end_c in components:
            top_c = table.cell(head_row1, start_c)
            top_c.merge(table.cell(head_row1, end_c))
            format_cell(top_c, comp_name, bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            sub_c1 = table.cell(head_row2, start_c)
            format_cell(sub_c1, "Actual\nTemp", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=7.5)

            sub_c2 = table.cell(head_row2, end_c)
            format_cell(sub_c2, "Temp\nRise", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=7.5)

        # Ambient Col
        c_amb_h = table.cell(head_row1, ambient_c)
        c_amb_h.merge(table.cell(head_row2, ambient_c))
        format_cell(c_amb_h, "Ambient", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        # Noise Col
        if noise_c is not None:
            c_noise_h = table.cell(head_row1, noise_c)
            c_noise_h.merge(table.cell(head_row2, noise_c))
            format_cell(c_noise_h, "Noise\n(dB)", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        # Vibration Col
        if vib_c is not None:
            c_vib_h = table.cell(head_row1, vib_c)
            c_vib_h.merge(table.cell(head_row2, vib_c))
            format_cell(c_vib_h, "Vibration\n(cm/s)", bold=True, bg_color="E2E8F0", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        current_row_idx += 2

        # 6. Data Rows
        threshold = 40.0
        if metadata.temp_rise_limit:
            m_th = re.search(r'\d+(?:\.\d+)?', metadata.temp_rise_limit)
            if m_th:
                threshold = float(m_th.group(0))

        act_keys = ['input_actual', 'body_actual', 'body2_actual', 'bc1_actual', 'bc2_actual', 'bc3_actual', 'bc4_actual', 'bc5_actual', 'output_actual']
        rise_keys = ['input_rise', 'body_rise', 'body2_rise', 'bc1_rise', 'bc2_rise', 'bc3_rise', 'bc4_rise', 'bc5_rise', 'output_rise']

        for row_idx, item in enumerate(intervals):
            r_bg = "F8FAFC" if row_idx % 2 == 1 else None

            # Time
            raw_time = str(item.time_label or '').strip()
            clean_time = re.split(r'\s*\(', raw_time)[0].strip() if '(' in raw_time else raw_time
            format_cell(table.cell(current_row_idx, 0), clean_time, bold=True, bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            # Direction (Strictly CW / CCW)
            raw_dir = str(getattr(item, 'direction', '') or '').strip().upper()
            clean_dir = 'CCW' if 'CCW' in raw_dir else 'CW'
            format_cell(table.cell(current_row_idx, 1), clean_dir, bold=True, bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            # Dynamic Channels
            for ch_i in range(num_channels):
                act_k = act_keys[ch_i] if ch_i < len(act_keys) else f"ch_{ch_i}_actual"
                rise_k = rise_keys[ch_i] if ch_i < len(rise_keys) else f"ch_{ch_i}_rise"
                act_v = getattr(item, act_k, 0.0) or 0.0
                rise_v = getattr(item, rise_k, 0.0) or 0.0

                format_cell(table.cell(current_row_idx, 2 + ch_i * 2), f"{act_v:.1f}", bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

                # Temp rise color
                rise_color = (0xDC, 0x26, 0x26) if rise_v > threshold else (0x02, 0x84, 0xC7)
                rise_str = f"+{rise_v:.1f}" if rise_v > 0 else f"{rise_v:.1f}"
                format_cell(table.cell(current_row_idx, 3 + ch_i * 2), rise_str, bold=True, color_rgb=rise_color, bg_color="F0F9FF" if rise_v <= threshold else "FEF2F2", align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            # Ambient
            amb_v = item.ambient
            format_cell(table.cell(current_row_idx, ambient_c), f"{amb_v:.1f}" if isinstance(amb_v, (int, float)) else str(amb_v), bold=True, color_rgb=(0x25, 0x63, 0xEB), bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            # Noise
            if noise_c is not None:
                n_val = getattr(item, 'noise', None)
                if n_val is None or n_val == 0.0:
                    n_val = metadata.noise_level_measured or "-"
                    try:
                        n_val = f"{float(str(n_val).replace('dB', '').strip()):.1f}"
                    except Exception:
                        pass
                else:
                    n_val = f"{float(n_val):.1f}"
                format_cell(table.cell(current_row_idx, noise_c), str(n_val), bold=True, color_rgb=(0x15, 0x80, 0x3D), bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            # Vibration
            if vib_c is not None:
                v_val = getattr(item, 'vibration', None)
                v_str = f"{float(v_val):.2f}" if v_val is not None else "-"
                format_cell(table.cell(current_row_idx, vib_c), v_str, bg_color=r_bg, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

            current_row_idx += 1

        # 7. Lubrication Footer
        c_l1 = table.cell(current_row_idx, 0)
        c_l1.merge(table.cell(current_row_idx, 1))
        format_cell(c_l1, "Lubrication leakage", bold=True, font_size=8.0)

        c_l2 = table.cell(current_row_idx, 2)
        c_l2.merge(table.cell(current_row_idx, total_cols - 1))
        format_cell(c_l2, metadata.lubrication_leakage or "No leakage", bold=True, color_rgb=(0x05, 0x96, 0x69), align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.0)

        out_path = get_output_path(output_filename)
        doc.save(str(out_path))
        return out_path
